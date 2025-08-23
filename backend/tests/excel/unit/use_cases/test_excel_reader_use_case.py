from io import BytesIO
from typing import cast

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend.features.excel.domain.use_cases.excel_reader_use_case import ExcelReaderUseCase


@pytest.fixture
def simple_excel_bytes() -> bytes:
    """Create a simple Excel file with basic data."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Sheet1"

    # Add headers
    ws.append(["Name", "Age", "City"])

    # Add data rows
    ws.append(["John Doe", 30, "New York"])
    ws.append(["Jane Smith", 25, "London"])
    ws.append(["Bob Johnson", 35, "Paris"])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def excel_with_empty_rows_bytes() -> bytes:
    """Create an Excel file with empty rows in between data."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Sheet1"

    # Add headers
    ws.append(["Product", "Price", "Stock"])

    # Add data with empty rows
    ws.append(["Apple", 1.5, 100])
    ws.append([None, None, None])
    ws.append(["", "", ""])
    ws.append(["Banana", 0.8, 50])
    ws.append([None, None, None])
    ws.append(["Orange", 2.0, 75])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def excel_with_multiple_sheets_bytes() -> bytes:
    """Create an Excel file with multiple sheets."""
    wb = Workbook()

    # First sheet
    ws1 = cast(Worksheet, wb.active)
    ws1.title = "Products"
    ws1.append(["Product", "Category"])
    ws1.append(["Laptop", "Electronics"])
    ws1.append(["Book", "Education"])

    # Second sheet
    ws2 = wb.create_sheet("Customers")
    ws2.append(["Customer", "Email"])
    ws2.append(["Alice", "alice@example.com"])
    ws2.append(["Bob", "bob@example.com"])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def excel_with_none_values_bytes() -> bytes:
    """Create an Excel file with None/null values."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Sheet1"

    # Add headers
    ws.append(["Name", None, "Score"])

    # Add data with mixed None values
    ws.append(["Alice", None, 95])
    ws.append([None, "Test", 88])
    ws.append(["Charlie", "Final", None])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def excel_custom_header_row_bytes() -> bytes:
    """Create an Excel file with headers not in the first row."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Sheet1"

    # Add some metadata rows
    ws.append(["Report Title", "Sales Data"])
    ws.append(["Generated On", "2025-01-01"])
    ws.append([])

    # Headers in row 4
    ws.append(["Item", "Quantity", "Revenue"])

    # Data rows
    ws.append(["Widget A", 10, 100.0])
    ws.append(["Widget B", 5, 75.5])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_read_simple_excel_file(simple_excel_bytes: bytes):
    """Test reading a simple Excel file with basic data."""
    reader = ExcelReaderUseCase(simple_excel_bytes)
    result = reader()

    expected = [
        {"Name": "John Doe", "Age": 30, "City": "New York"},
        {"Name": "Jane Smith", "Age": 25, "City": "London"},
        {"Name": "Bob Johnson", "Age": 35, "City": "Paris"},
    ]

    assert result == expected
    assert len(result) == 3


def test_read_excel_with_empty_rows(excel_with_empty_rows_bytes: bytes):
    """Test that empty rows are properly skipped."""
    reader = ExcelReaderUseCase(excel_with_empty_rows_bytes)
    result = reader()

    expected = [
        {"Product": "Apple", "Price": 1.5, "Stock": 100},
        {"Product": "Banana", "Price": 0.8, "Stock": 50},
        {"Product": "Orange", "Price": 2.0, "Stock": 75},
    ]

    assert result == expected
    assert len(result) == 3


def test_read_specific_sheet(excel_with_multiple_sheets_bytes: bytes):
    """Test reading from a specific sheet."""
    reader = ExcelReaderUseCase(excel_with_multiple_sheets_bytes)

    # Read from default sheet (Products)
    result_default = reader()
    expected_products = [
        {"Product": "Laptop", "Category": "Electronics"},
        {"Product": "Book", "Category": "Education"},
    ]
    assert result_default == expected_products

    # Read from specific sheet (Customers)
    result_customers = reader(sheet="Customers")
    expected_customers = [
        {"Customer": "Alice", "Email": "alice@example.com"},
        {"Customer": "Bob", "Email": "bob@example.com"},
    ]
    assert result_customers == expected_customers


def test_read_with_none_values(excel_with_none_values_bytes: bytes):
    """Test handling of None values in headers and data."""
    reader = ExcelReaderUseCase(excel_with_none_values_bytes)
    result = reader()

    expected = [
        {"Name": "Alice", "": None, "Score": 95},
        {"Name": None, "": "Test", "Score": 88},
        {"Name": "Charlie", "": "Final", "Score": None},
    ]

    assert result == expected
    assert len(result) == 3


def test_custom_header_row(excel_custom_header_row_bytes: bytes):
    """Test reading with custom header row position."""
    reader = ExcelReaderUseCase(excel_custom_header_row_bytes, header_row=4)
    result = reader()

    expected = [
        {"Item": "Widget A", "Quantity": 10, "Revenue": 100.0},
        {"Item": "Widget B", "Quantity": 5, "Revenue": 75.5},
    ]

    assert result == expected
    assert len(result) == 2


@pytest.mark.parametrize(
    "header_row,expected_length",
    [
        (1, 3),  # Normal case
        (2, 2),  # Skip first data row
        (3, 1),  # Skip first two data rows
    ],
)
def test_different_header_row_positions(
    simple_excel_bytes: bytes, header_row: int, expected_length: int
):
    """Test reading with different header row positions."""
    reader = ExcelReaderUseCase(simple_excel_bytes, header_row=header_row)
    result = reader()

    assert len(result) == expected_length
    assert isinstance(result, list)
    if result:  # If there are results
        assert isinstance(result[0], dict)


def test_nonexistent_sheet_raises_error(simple_excel_bytes: bytes):
    """Test that requesting a non-existent sheet raises an error."""
    reader = ExcelReaderUseCase(simple_excel_bytes)

    with pytest.raises(KeyError, match="Worksheet NonExistentSheet does not exist."):
        reader(sheet="NonExistentSheet")


def test_empty_excel_file():
    """Test reading an Excel file with only headers."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    ws.append(["Header1", "Header2", "Header3"])
    # No data rows

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    reader = ExcelReaderUseCase(excel_bytes)
    result = reader()

    assert result == []
    assert len(result) == 0


@pytest.mark.parametrize("sheet_name", ["Products", "Customers", None])  # Default sheet
def test_sheet_parameter_variations(excel_with_multiple_sheets_bytes: bytes, sheet_name: str):
    """Test different sheet name parameters."""
    reader = ExcelReaderUseCase(excel_with_multiple_sheets_bytes)

    if sheet_name is None:
        result = reader()
    else:
        result = reader(sheet=sheet_name)

    assert isinstance(result, list)
    if result:
        assert isinstance(result[0], dict)


def test_whitespace_handling_in_headers():
    """Test that whitespace in headers is properly stripped."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)

    # Add headers with whitespace
    ws.append([" Name ", "  Age", "City  "])
    ws.append(["John", 30, "NYC"])

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    reader = ExcelReaderUseCase(excel_bytes)
    result = reader()

    expected = [{"Name": "John", "Age": 30, "City": "NYC"}]
    assert result == expected


def test_mixed_data_types():
    """Test reading Excel file with mixed data types."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)

    ws.append(["Text", "Number", "Float", "Boolean"])
    ws.append(["Hello", 42, 3.14, True])
    ws.append(["World", 0, 2.71, False])

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    reader = ExcelReaderUseCase(excel_bytes)
    result = reader()

    expected = [
        {"Text": "Hello", "Number": 42, "Float": 3.14, "Boolean": True},
        {"Text": "World", "Number": 0, "Float": 2.71, "Boolean": False},
    ]

    assert result == expected
