from io import BytesIO
from typing import Optional

from openpyxl import load_workbook


class ExcelReaderUseCase:
    def __init__(self, file_bytes: bytes, *, header_row: int = 1):
        self.file_bytes = file_bytes
        self.header_row = header_row

    def __call__(self, *, sheet: Optional[str] = None) -> list[dict[str, object]]:
        wb = load_workbook(BytesIO(self.file_bytes), data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb.active

        hdr = next(ws.iter_rows(min_row=self.header_row, max_row=self.header_row, values_only=True))
        headers = [str(h).strip() if h is not None else "" for h in hdr]

        rows: list[dict[str, object]] = []
        for row in ws.iter_rows(min_row=self.header_row + 1, values_only=True):
            if all(v in (None, "") for v in row):
                continue
            rows.append({headers[i]: row[i] for i in range(len(headers))})
        return rows
